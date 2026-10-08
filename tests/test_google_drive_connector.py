import asyncio
import io
import threading
import zipfile
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime

import httpx
import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference

from platform_be.services.connectors import google_drive, header_rows
from platform_be.services.connectors.base import (
    REASON_MESSAGES,
    Column,
    ConnectorError,
    QuerySource,
    TableRef,
    TableSource,
)
from platform_be.services.connectors.google_drive import (
    CSV,
    NO_FILE,
    NO_FOLDER,
    NOT_CSV,
    NOT_XLSX,
    SAME_NAME,
    SPREADSHEET,
    XLSX,
    GoogleDriveConnector,
    parse_folder_id,
    quoted,
)
from platform_be.services.connectors.values import to_text
from tests.fakes import (
    FakeDriveFile,
    FakeGoogleDrive,
    FakeGoogleDriveOAuth,
    FakeSpreadsheet,
    access_token_for,
)

FOLDER_ID = "1AbCdEfGhIjKlMnOpQrStUvWxYz012345"
OTHER_FOLDER_ID = "1ZyXwVuTsRqPoNmLkJiHgFeDcBa987654"
SHEET_ID = "1BxiMVs0XRA5nFMdKvBdBZjgmUUqptlbs74OgvE2upms"
REFRESH_TOKEN = "refresh-token-of-the-owner"
MAX_FILE_BYTES = 20_000
SCORES_CSV = b'student,school,score\r\nAn,A,7.5\r\nBinh,"B, north",8\r\nChi,,6.25\r\n'
SCORES = [("An", "A", "7.5"), ("Binh", "B, north", "8"), ("Chi", None, "6.25")]
TEXT = [Column("student", "text"), Column("school", "text"), Column("score", "text")]


def workbook() -> bytes:
    """A workbook as Excel saves one: two sheets of cells and a sheet that is a chart."""
    book = openpyxl.Workbook()
    scores = book.active
    scores.title = "Scores"
    scores.append(["student", "score", "taken", "day", "passed", "double", "group"])
    scores.append(["An", 7.5, datetime(2026, 3, 1, 9, 30), date(2026, 3, 1), True, "=B2*2", "x"])
    scores.append(["Binh", 8, None, None, False, None, None])
    # Only the first cell of a merged range holds its value.
    scores.merge_cells("G2:G3")
    book.create_sheet("Second").append(["only"])
    book["Second"].append([1])
    chart = BarChart()
    chart.add_data(Reference(scores, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    book.create_chartsheet("Chart").add_chart(chart)
    saved = io.BytesIO()
    book.save(saved)
    return saved.getvalue()


def folder(**files: FakeDriveFile) -> tuple[FakeGoogleDrive, FakeGoogleDriveOAuth]:
    """A Drive with the pinned folder, another folder, and the files given by ID."""
    drive = FakeGoogleDrive()
    drive.readers = {REFRESH_TOKEN}
    drive.files = {
        FOLDER_ID: FakeDriveFile("Survey data", FakeGoogleDrive.FOLDER, {"root"}),
        OTHER_FOLDER_ID: FakeDriveFile("Private", FakeGoogleDrive.FOLDER, {"root"}),
        **files,
    }
    return drive, FakeGoogleDriveOAuth()


def csv_file(name: str, content: bytes = SCORES_CSV, **changes: object) -> FakeDriveFile:
    return FakeDriveFile(
        **{"name": name, "mime_type": CSV, "parents": {FOLDER_ID}, "content": content} | changes
    )


@pytest.fixture
def executor():
    with ThreadPoolExecutor(max_workers=2) as threads:
        yield threads


@pytest.fixture
def temporary_files(monkeypatch) -> list:
    """Every temporary file the connector opens."""
    opened = []
    make = google_drive.tempfile.TemporaryFile

    def recording():
        opened.append(make())
        return opened[-1]

    monkeypatch.setattr(google_drive.tempfile, "TemporaryFile", recording)
    return opened


def connector(
    drive: FakeGoogleDrive,
    oauth: FakeGoogleDriveOAuth,
    executor: ThreadPoolExecutor,
    *,
    folder_id: str = FOLDER_ID,
    refresh_token: str = REFRESH_TOKEN,
) -> GoogleDriveConnector:
    return GoogleDriveConnector(
        {"folder_id": folder_id},
        {"refresh_token": refresh_token},
        oauth=oauth,
        executor=executor,
        max_file_bytes=MAX_FILE_BYTES,
        connect_timeout=1,
        query_timeout=1,
        transport=drive.transport,
    )


def table(schema: str, name: str | None = None) -> TableSource:
    return TableSource(type="table", schema=schema, name=name or schema)


async def read(reader: GoogleDriveConnector, source, max_rows: int | None = None):
    """The columns and the rows of a source, the values as the text a dataset stores."""
    async with reader.open_rows(source, max_rows=max_rows) as stream:
        rows = [tuple(to_text(value) for value in row) async for row in stream.rows]
        return stream.columns, rows


async def reason(reading) -> tuple[str, str]:
    with pytest.raises(ConnectorError) as raised:
        await reading
    return raised.value.reason, raised.value.message


@pytest.mark.parametrize(
    "address",
    [
        FOLDER_ID,
        f"  {FOLDER_ID} ",
        f"https://drive.google.com/drive/folders/{FOLDER_ID}",
        f"https://drive.google.com/drive/folders/{FOLDER_ID}?usp=sharing",
        f"https://drive.google.com/drive/u/0/folders/{FOLDER_ID}",
        f"https://drive.google.com/drive/u/12/folders/{FOLDER_ID}/",
    ],
)
def test_a_folder_is_named_by_its_address_or_its_id(address: str) -> None:
    assert parse_folder_id(address) == FOLDER_ID


@pytest.mark.parametrize(
    "address",
    [
        "",
        "short",
        "has spaces in the middle",
        f"http://drive.google.com/drive/folders/{FOLDER_ID}",
        f"https://drive.google.com.evil.example/drive/folders/{FOLDER_ID}",
        f"https://evil.example/?https://drive.google.com/drive/folders/{FOLDER_ID}",
        f"https://drive.google.com/file/d/{FOLDER_ID}/view",
        f"https://docs.google.com/spreadsheets/d/{FOLDER_ID}/edit",
        f"{FOLDER_ID}' or name contains '",
    ],
)
def test_anything_else_is_not_a_folder(address: str) -> None:
    with pytest.raises(ValueError):
        parse_folder_id(address)


def test_a_quote_or_a_backslash_cannot_end_a_string_of_a_search() -> None:
    assert quoted("plain.csv") == "'plain.csv'"
    assert quoted("it's") == r"'it\'s'"
    assert quoted("a\\b") == r"'a\\b'"
    # The backslash first: the one that escapes a quote is not escaped again.
    assert quoted("\\'") == r"'\\\''"


async def test_the_test_reads_the_folder_and_keeps_its_name(executor) -> None:
    drive, oauth = folder(**{"file-1": csv_file("scores.csv")})
    reader = connector(drive, oauth, executor)
    await reader.test()
    assert reader.folder_name == "Survey data"
    assert oauth.refreshed == [REFRESH_TOKEN]

    # A file is not a folder, whoever can open it.
    not_a_folder = connector(drive, oauth, executor, folder_id="file-1")
    assert await reason(not_a_folder.test()) == ("permission_denied", NO_FOLDER)
    unknown = connector(drive, oauth, executor, folder_id="no-such-folder-anywhere")
    assert await reason(unknown.test()) == ("permission_denied", NO_FOLDER)
    stranger = connector(drive, oauth, executor, refresh_token="somebody-else")
    assert await reason(stranger.test()) == ("permission_denied", NO_FOLDER)

    oauth.revoked.add(REFRESH_TOKEN)
    assert (await reason(connector(drive, oauth, executor).test()))[0] == "access_revoked"


async def test_the_readable_files_directly_in_the_folder_are_the_schemas(executor) -> None:
    drive, oauth = folder(
        **{
            "f-csv": csv_file("b scores.csv"),
            "f-xlsx": FakeDriveFile("a book.xlsx", XLSX, {FOLDER_ID}),
            SHEET_ID: FakeDriveFile("c sheet", SPREADSHEET, {FOLDER_ID}),
            "f-pdf": FakeDriveFile("report.pdf", "application/pdf", {FOLDER_ID}),
            "f-doc": FakeDriveFile("notes", "application/vnd.google-apps.document", {FOLDER_ID}),
            "f-xls": FakeDriveFile("old.xls", "application/vnd.ms-excel", {FOLDER_ID}),
            "f-sub": FakeDriveFile("inner", FakeGoogleDrive.FOLDER, {FOLDER_ID}),
            "f-deep": csv_file("deep.csv", parents={"f-sub"}),
            "f-gone": csv_file("binned.csv", trashed=True),
            "f-else": csv_file("elsewhere.csv", parents={OTHER_FOLDER_ID}),
        }
    )
    reader = connector(drive, oauth, executor)
    assert await reader.list_schemas() == ["a book.xlsx", "b scores.csv", "c sheet"]
    assert drive.downloads == []


async def test_a_long_folder_is_listed_page_by_page_up_to_a_limit(executor, monkeypatch) -> None:
    drive, oauth = folder(**{f"f-{n}": csv_file(f"file {n}.csv") for n in range(7)})
    drive.page_size = 3
    reader = connector(drive, oauth, executor)
    assert await reader.list_schemas() == [f"file {n}.csv" for n in range(7)]
    assert len(drive.queries) == 3

    drive.queries.clear()
    monkeypatch.setattr(google_drive, "MAX_FILES", 4)
    assert await reader.list_schemas() == [f"file {n}.csv" for n in range(4)]
    assert len(drive.queries) == 2


async def test_a_csv_file_is_one_table_named_like_the_file(executor, temporary_files) -> None:
    drive, oauth = folder(**{"f-csv": csv_file("scores.csv")})
    reader = connector(drive, oauth, executor)

    assert await reader.list_tables("scores.csv", search=None, limit=10) == [
        TableRef(schema="scores.csv", name="scores.csv", type="table", column_count=None)
    ]
    assert await reader.list_tables("scores.csv", search="nothing", limit=10) == []
    assert await reader.list_tables("gone.csv", search=None, limit=10) == []
    assert await reader.list_columns("scores.csv", "scores.csv") == TEXT
    assert await read(reader, table("scores.csv")) == (TEXT, SCORES)
    assert await read(reader, table("scores.csv"), max_rows=2) == (TEXT, SCORES[:2])
    assert await read(reader, table("scores.csv"), max_rows=0) == (TEXT, [])

    assert await reason(reader.list_columns("scores.csv", "other")) == (
        "source_not_found",
        google_drive.NO_TAB,
    )
    assert await reason(read(reader, table("gone.csv"))) == ("source_not_found", NO_FILE)
    assert (await reason(read(reader, QuerySource(type="query", sql="SELECT 1"))))[0] == (
        "unsupported_source"
    )
    # One token for all of it, and no file left behind.
    assert oauth.refreshed == [REFRESH_TOKEN]
    assert temporary_files and all(handle.closed for handle in temporary_files)


async def test_a_csv_file_is_read_as_utf8_text_with_a_header_row(executor, temporary_files) -> None:
    drive, oauth = folder(
        **{
            "f-bom": csv_file("excel.csv", '﻿name,note\r\nÂn,"two\nlines"\r\n'.encode()),
            "f-latin": csv_file("latin.csv", "name\ncafé\n".encode("latin-1")),
            "f-empty": csv_file("empty.csv", b""),
            "f-header": csv_file("header.csv", b"a,b\n"),
            # Empty lines are rows only when a row with something in it follows.
            "f-gaps": csv_file("gaps.csv", b"a,b\n1,2\n,\n\n3,\n\n,\n"),
            "f-same": csv_file("same.csv", b"a,a\n1,2\n"),
            "f-stray": csv_file("stray.csv", b"a,,c\n1,,3\n1,x,3\n"),
            "f-long": csv_file("long.csv", b"a\n1,2\n"),
            "f-nul": csv_file("binary.csv", b"a,b\n1,\x00\n"),
        }
    )
    reader = connector(drive, oauth, executor)
    named = [Column("name", "text"), Column("note", "text")]
    assert await read(reader, table("excel.csv")) == (named, [("Ân", "two\nlines")])
    assert await reason(read(reader, table("latin.csv"))) == ("source_malformed", NOT_CSV)
    assert await reader.list_columns("empty.csv", "empty.csv") == []
    assert await read(reader, table("empty.csv")) == ([], [])
    ab = [Column("a", "text"), Column("b", "text")]
    assert await read(reader, table("header.csv")) == (ab, [])
    gaps = [("1", "2"), (None, None), (None, None), ("3", None)]
    assert await read(reader, table("gaps.csv")) == (ab, gaps)
    assert await read(reader, table("gaps.csv"), max_rows=3) == (ab, gaps[:3])
    default = REASON_MESSAGES["source_malformed"]
    assert await reason(reader.list_columns("same.csv", "same.csv")) == (
        "source_malformed",
        default,
    )
    assert await reason(read(reader, table("stray.csv"))) == ("source_malformed", default)
    assert await reason(read(reader, table("long.csv"))) == ("source_malformed", default)
    # No text column can store a NUL character.
    assert await reason(read(reader, table("binary.csv"))) == ("source_malformed", NOT_CSV)
    assert all(handle.closed for handle in temporary_files)


async def test_many_rows_are_handed_over_a_batch_at_a_time(executor, monkeypatch) -> None:
    monkeypatch.setattr(google_drive, "ROWS_PER_BATCH", 3)
    lines = b"n\n" + b"".join(b"%d\n" % n for n in range(10))
    drive, oauth = folder(**{"f-csv": csv_file("numbers.csv", lines)})
    reader = connector(drive, oauth, executor)
    _, rows = await read(reader, table("numbers.csv"))
    assert rows == [(str(n),) for n in range(10)]
    _, rows = await read(reader, table("numbers.csv"), max_rows=7)
    assert rows == [(str(n),) for n in range(7)]


async def test_a_name_is_looked_up_exactly_and_only_in_the_folder(executor) -> None:
    drive, oauth = folder(
        **{
            "f-quote": csv_file("Q1's data.csv", b"a\nquote\n"),
            "f-slash": csv_file("back\\slash.csv", b"a\nslash\n"),
            "f-trick": csv_file("x' or name contains '", b"a\ntrick\n"),
            # The same account opens these, and the connection does not.
            "f-else": csv_file("elsewhere.csv", b"secret\n1\n", parents={OTHER_FOLDER_ID}),
            "f-deep": csv_file("deep.csv", b"secret\n1\n", parents={"f-sub"}),
            "f-sub": FakeDriveFile("inner", FakeGoogleDrive.FOLDER, {FOLDER_ID}),
            "f-gone": csv_file("binned.csv", b"secret\n1\n", trashed=True),
            "f-pdf": FakeDriveFile("report.pdf", "application/pdf", {FOLDER_ID}, b"%PDF"),
        }
    )
    reader = connector(drive, oauth, executor)
    for name, value in (
        ("Q1's data.csv", "quote"),
        ("back\\slash.csv", "slash"),
        ("x' or name contains '", "trick"),
    ):
        assert await read(reader, table(name)) == ([Column("a", "text")], [(value,)]), name

    for name in ("elsewhere.csv", "deep.csv", "binned.csv", "report.pdf", "inner", "f-else"):
        assert await reason(read(reader, table(name))) == ("source_not_found", NO_FILE), name
        assert await reason(reader.list_columns(name, name)) == ("source_not_found", NO_FILE)
        assert await reader.list_tables(name, search=None, limit=10) == []
    assert drive.downloads == ["f-quote", "f-slash", "f-trick"]
    assert drive.queries
    assert all(query.startswith(f"'{FOLDER_ID}' in parents and ") for query in drive.queries)


async def test_only_the_file_with_exactly_the_name_asked_for_is_read(executor) -> None:
    drive, oauth = folder(
        **{
            "f-lower": csv_file("scores.csv", b"a\nlower\n"),
            "f-upper": csv_file("Scores.csv", b"a\nupper\n"),
            "f-pdf": FakeDriveFile("report.pdf", "application/pdf", {FOLDER_ID}, b"a\n1\n"),
        }
    )
    # A search that finds more than was asked for: what it finds is checked again.
    drive.loose_search = True
    reader = connector(drive, oauth, executor)
    assert await read(reader, table("scores.csv")) == ([Column("a", "text")], [("lower",)])
    assert await read(reader, table("Scores.csv")) == ([Column("a", "text")], [("upper",)])
    assert await reason(read(reader, table("SCORES.csv"))) == ("source_not_found", NO_FILE)
    assert await reason(read(reader, table("report.pdf"))) == ("source_not_found", NO_FILE)


async def test_a_name_two_files_share_is_not_read(executor) -> None:
    drive, oauth = folder(
        **{
            "f-1": csv_file("scores.csv", b"a\nfirst\n"),
            "f-2": FakeDriveFile("scores.csv", XLSX, {FOLDER_ID}, workbook()),
            "f-3": csv_file("alone.csv", b"a\n1\n"),
            # Not a second one: it is in another folder.
            "f-4": csv_file("alone.csv", b"a\n2\n", parents={OTHER_FOLDER_ID}),
        }
    )
    reader = connector(drive, oauth, executor)
    assert await reader.list_schemas() == ["alone.csv", "scores.csv"]
    assert await read(reader, table("alone.csv")) == ([Column("a", "text")], [("1",)])
    for reading in (
        read(reader, table("scores.csv")),
        reader.list_columns("scores.csv", "scores.csv"),
        reader.list_tables("scores.csv", search=None, limit=10),
    ):
        assert await reason(reading) == ("source_malformed", SAME_NAME)
    assert drive.downloads == ["f-3"]


async def test_a_file_larger_than_a_dataset_may_be_is_not_fetched(
    executor, temporary_files
) -> None:
    big = b"a\n" + b"x" * MAX_FILE_BYTES
    drive, oauth = folder(
        **{
            "f-big": csv_file("big.csv", big),
            "f-unsized": csv_file("unsized.csv", big, sized=False),
            "f-fits": csv_file("fits.csv", b"a\n" + b"x" * (MAX_FILE_BYTES - 2)),
        }
    )
    reader = connector(drive, oauth, executor)
    # Drive says how large it is: refused before a byte is asked for.
    assert (await reason(read(reader, table("big.csv"))))[0] == "source_too_large"
    assert drive.downloads == [] and temporary_files == []
    # Drive does not say: cut off while it arrives.
    assert (await reason(read(reader, table("unsized.csv"))))[0] == "source_too_large"
    assert (await reason(reader.list_columns("unsized.csv", "unsized.csv")))[0] == (
        "source_too_large"
    )
    assert drive.downloads == ["f-unsized", "f-unsized"]
    _, rows = await read(reader, table("fits.csv"))
    assert len(rows) == 1
    assert len(temporary_files) == 3 and all(handle.closed for handle in temporary_files)


async def test_an_excel_workbook_has_a_table_for_each_sheet_of_cells(
    executor, temporary_files
) -> None:
    drive, oauth = folder(**{"f-xlsx": FakeDriveFile("book.xlsx", XLSX, {FOLDER_ID}, workbook())})
    reader = connector(drive, oauth, executor)

    # The chart is not a table.
    assert await reader.list_tables("book.xlsx", search=None, limit=10) == [
        TableRef(schema="book.xlsx", name=name, type="table", column_count=None)
        for name in ["Scores", "Second"]
    ]
    assert await reader.list_tables("book.xlsx", search="SEC", limit=10) == [
        TableRef(schema="book.xlsx", name="Second", type="table", column_count=None)
    ]
    names = ["student", "score", "taken", "day", "passed", "double", "group"]
    columns = [Column(name, "text") for name in names]
    assert await reader.list_columns("book.xlsx", "Scores") == columns
    assert await read(reader, table("book.xlsx", "Scores")) == (
        columns,
        [
            # A formula that no program has computed has no value; a merged range has one.
            ("An", "7.5", "2026-03-01T09:30:00", "2026-03-01T00:00:00", "true", None, "x"),
            ("Binh", "8", None, None, "false", None, None),
        ],
    )
    assert await read(reader, table("book.xlsx", "Second")) == ([Column("only", "text")], [("1",)])
    assert await read(reader, table("book.xlsx", "Scores"), max_rows=1) == (
        columns,
        [("An", "7.5", "2026-03-01T09:30:00", "2026-03-01T00:00:00", "true", None, "x")],
    )
    for missing in ("Chart", "Gone"):
        assert await reason(read(reader, table("book.xlsx", missing))) == (
            "source_not_found",
            google_drive.NO_TAB,
        )
    assert temporary_files and all(handle.closed for handle in temporary_files)


async def test_a_file_that_is_not_a_workbook_is_refused(executor, temporary_files) -> None:
    not_a_workbook = io.BytesIO()
    with zipfile.ZipFile(not_a_workbook, "w") as archive:
        archive.writestr("readme.txt", "not a workbook")
    damaged = workbook()
    drive, oauth = folder(
        **{
            "f-text": FakeDriveFile("text.xlsx", XLSX, {FOLDER_ID}, b"just some text"),
            "f-zip": FakeDriveFile("zip.xlsx", XLSX, {FOLDER_ID}, not_a_workbook.getvalue()),
            "f-cut": FakeDriveFile("cut.xlsx", XLSX, {FOLDER_ID}, damaged[: len(damaged) // 2]),
            "f-none": FakeDriveFile("none.xlsx", XLSX, {FOLDER_ID}, b""),
        }
    )
    reader = connector(drive, oauth, executor)
    for name in ("text.xlsx", "zip.xlsx", "cut.xlsx", "none.xlsx"):
        for reading in (
            reader.list_tables(name, search=None, limit=10),
            reader.list_columns(name, "Sheet"),
            read(reader, table(name, "Sheet")),
        ):
            assert await reason(reading) == ("source_malformed", NOT_XLSX), name
    assert len(temporary_files) == 12 and all(handle.closed for handle in temporary_files)


SHEET_XML = (
    '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">%s</worksheet>'
)
TEXTS_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"


SHEET1 = "xl/worksheets/sheet1.xml"
TEXTS = "xl/sharedStrings.xml"


def crafted(parts: dict[str, bytes], types: dict[str, str] | None = None) -> bytes:
    """A real workbook with some of its parts replaced or added. `types` says what the file
    claims a part is."""
    parts = dict(parts)
    saved = io.BytesIO()
    book = openpyxl.Workbook()
    book.active.append(["a"])
    book.active.append(["kept"])
    book.save(saved)
    made = io.BytesIO()
    with zipfile.ZipFile(saved) as source, zipfile.ZipFile(made, "w", zipfile.ZIP_DEFLATED) as out:
        for entry in source.infolist():
            content = parts.pop(entry.filename, None) or source.read(entry.filename)
            if entry.filename == "[Content_Types].xml":
                claims = "".join(
                    f'<Override PartName="/{name}" ContentType="{kind}"/>'
                    for name, kind in (types or {}).items()
                )
                content = content.replace(b"</Types>", claims.encode() + b"</Types>")
            out.writestr(entry.filename, content)
        for name, content in parts.items():
            out.writestr(name, content)
    return made.getvalue()


def xlsx_folder(**files: bytes) -> tuple[FakeGoogleDrive, FakeGoogleDriveOAuth]:
    return folder(
        **{
            f"f-{name}": FakeDriveFile(f"{name}.xlsx", XLSX, {FOLDER_ID}, content)
            for name, content in files.items()
        }
    )


async def test_a_workbook_that_unpacks_to_too_much_is_not_opened(
    executor, temporary_files, monkeypatch
) -> None:
    limit = MAX_FILE_BYTES * google_drive.MAX_UNPACKED_RATIO
    half = b"0" * (limit // 2)
    bomb = crafted({SHEET1: half, "xl/worksheets/sheet2.xml": half})
    assert len(bomb) < MAX_FILE_BYTES
    drive, oauth = xlsx_folder(bomb=bomb)
    reader = connector(drive, oauth, executor)

    def never(*args, **kwargs):
        raise AssertionError("the workbook was opened")

    monkeypatch.setattr(google_drive._WorkbookReader, "read", never)
    for reading in (
        reader.list_tables("bomb.xlsx", search=None, limit=10),
        read(reader, table("bomb.xlsx", "Sheet")),
    ):
        assert (await reason(reading))[0] == "source_too_large"
    assert drive.downloads == ["f-bomb", "f-bomb"]
    assert all(handle.closed for handle in temporary_files)


async def test_no_part_of_a_workbook_is_held_in_memory_past_its_limit(
    executor, temporary_files, monkeypatch
) -> None:
    monkeypatch.setattr(google_drive, "MAX_PART_BYTES", 20_000)
    texts_limit = MAX_FILE_BYTES * google_drive.MAX_TEXTS_RATIO
    assert texts_limit > 30_000 > google_drive.MAX_PART_BYTES

    def texts(size: int) -> bytes:
        ns = b'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
        return b"<sst " + ns + b"><si><t>kept</t></si>" + b" " * size + b"</sst>"

    cell = b'<sheetData><row r="1"><c r="A1" t="s"><v>0</v></c></row></sheetData>'
    shared = {TEXTS: TEXTS_TYPE}
    moved = {"xl/worksheets/sheet7.xml": TEXTS_TYPE}
    padding = b"<!--" + b"0" * 30_000 + b"-->"
    drive, oauth = xlsx_folder(
        # The table of texts has more room than the other parts, and no more than its own.
        texts=crafted({TEXTS: texts(30_000), SHEET1: SHEET_XML.encode() % cell}, shared),
        toomany=crafted({TEXTS: texts(texts_limit)}, shared),
        # Wherever the file says the table is, and whatever else it calls a part.
        moved=crafted({"xl/worksheets/sheet7.xml": texts(texts_limit)}, moved),
        elsewhere=crafted({"xl/worksheets/sheet7.xml": texts(30_000)}, moved),
        # The room is for reading it as the table of texts, not for whatever reads it next.
        twice=crafted({"xl/styles.xml": texts(30_000)}, {"xl/styles.xml": TEXTS_TYPE}),
        styles=crafted({"xl/styles.xml": b"<styleSheet>" + b"<xf/>" * 6_000 + b"</styleSheet>"}),
        links=crafted({"xl/_rels/workbook.xml.rels": padding}),
        # A sheet is read a row at a time: it may be larger than any of them.
        sheet=crafted({SHEET1: SHEET_XML.encode() % (padding + cell)}),
    )
    reader = connector(drive, oauth, executor)
    assert await read(reader, table("texts.xlsx", "Sheet")) == ([Column("kept", "text")], [])
    assert await read(reader, table("elsewhere.xlsx", "Sheet")) == (
        [Column("a", "text")],
        [("kept",)],
    )
    # Its cell points at a text the workbook does not have: read, and found broken.
    assert await reason(read(reader, table("sheet.xlsx", "Sheet"))) == (
        "source_malformed",
        NOT_XLSX,
    )
    assert await reader.list_tables("sheet.xlsx", search=None, limit=1) == [
        TableRef(schema="sheet.xlsx", name="Sheet", type="table", column_count=None)
    ]
    for name in ("toomany", "moved", "twice", "styles", "links"):
        for reading in (
            reader.list_tables(f"{name}.xlsx", search=None, limit=10),
            read(reader, table(f"{name}.xlsx", "Sheet")),
        ):
            assert (await reason(reading))[0] == "source_too_large", name
    assert all(handle.closed for handle in temporary_files)


async def test_a_run_of_empty_rows_does_not_go_on_for_ever(executor, monkeypatch) -> None:
    monkeypatch.setattr(header_rows, "MAX_EMPTY_ROWS", 50)

    def rows_at(*numbers: int) -> bytes:
        rows = "".join(
            f'<row r="{n}"><c r="A{n}" t="inlineStr"><is><t>v{n}</t></is></c></row>'
            for n in numbers
        )
        return (SHEET_XML % f"<sheetData>{rows}</sheetData>").encode()

    drive, oauth = folder(
        **{
            # A few bytes that claim a row far below the last one.
            "f-far": FakeDriveFile(
                "far.xlsx",
                XLSX,
                {FOLDER_ID},
                crafted({SHEET1: rows_at(1, 2, 99_999_999_999)}),
            ),
            "f-first": FakeDriveFile(
                "first.xlsx",
                XLSX,
                {FOLDER_ID},
                crafted({SHEET1: rows_at(99_999_999_999)}),
            ),
            "f-near": FakeDriveFile(
                "near.xlsx",
                XLSX,
                {FOLDER_ID},
                crafted({SHEET1: rows_at(1, 2, 53)}),
            ),
            "f-lines": csv_file("lines.csv", b"a\n1\n" + b"\n" * 51 + b"2\n"),
            "f-fewer": csv_file("fewer.csv", b"a\n1\n" + b"\n" * 50 + b"2\n"),
        }
    )
    reader = connector(drive, oauth, executor)
    for name, tab in (("far.xlsx", "Sheet"), ("first.xlsx", "Sheet"), ("lines.csv", "lines.csv")):
        assert await reason(read(reader, table(name, tab))) == (
            "source_malformed",
            header_rows.TOO_MANY_EMPTY_ROWS,
        ), name
    # What was read before the run is not lost, and a preview does not reach the run at all.
    assert await read(reader, table("far.xlsx", "Sheet"), max_rows=1) == (
        [Column("v1", "text")],
        [("v2",)],
    )
    _, rows = await read(reader, table("near.xlsx", "Sheet"))
    assert rows == [("v2",), *[(None,)] * 50, ("v53",)]
    _, rows = await read(reader, table("fewer.csv"))
    assert rows == [("1",), *[(None,)] * 50, ("2",)]


async def test_reading_a_file_does_not_hold_up_the_event_loop(executor, monkeypatch) -> None:
    loop_thread = threading.current_thread()
    seen = []
    names, rows = google_drive.column_names, google_drive.data_rows

    def watched_names(first_row):
        seen.append(threading.current_thread())
        return names(first_row)

    def watched_rows(cells, *args):
        for row in rows(cells, *args):
            seen.append(threading.current_thread())
            yield row

    monkeypatch.setattr(google_drive, "column_names", watched_names)
    monkeypatch.setattr(google_drive, "data_rows", watched_rows)
    opening = google_drive.open_workbook

    def watched_opening(*args):
        seen.append(threading.current_thread())
        return opening(*args)

    monkeypatch.setattr(google_drive, "open_workbook", watched_opening)
    drive, oauth = folder(
        **{
            "f-csv": csv_file("scores.csv"),
            "f-xlsx": FakeDriveFile("book.xlsx", XLSX, {FOLDER_ID}, workbook()),
        }
    )
    reader = connector(drive, oauth, executor)
    await read(reader, table("scores.csv"))
    await read(reader, table("book.xlsx", "Scores"))
    assert len(seen) >= 8 and loop_thread not in seen


async def test_a_spreadsheet_in_the_folder_is_read_through_the_sheets_api(executor) -> None:
    drive, oauth = folder(
        **{
            SHEET_ID: FakeDriveFile("Survey 2026", SPREADSHEET, {FOLDER_ID}),
            "1-a-spreadsheet-of-somebody-else-entirely": FakeDriveFile(
                "Private", SPREADSHEET, {OTHER_FOLDER_ID}
            ),
        }
    )
    drive.sheets.spreadsheets[SHEET_ID] = FakeSpreadsheet(
        # Renamed since: the file's name in the folder is what names the schema.
        title="Survey 2026 (renamed)",
        tabs={"Answers": [["student", "score"], ["An", 7.5], ["Binh", 8]], "Empty": []},
        readers={REFRESH_TOKEN},
        chart_tabs=["Chart1"],
    )
    drive.sheets.spreadsheets["1-a-spreadsheet-of-somebody-else-entirely"] = FakeSpreadsheet(
        title="Private", tabs={"Secret": [["a"], [1]]}, readers={REFRESH_TOKEN}
    )
    reader = connector(drive, oauth, executor)
    assert await reader.list_schemas() == ["Survey 2026"]
    assert await reader.list_tables("Survey 2026", search=None, limit=10) == [
        TableRef(schema="Survey 2026", name=name, type="table", column_count=None)
        for name in ["Answers", "Empty"]
    ]
    columns = [Column("student", "text"), Column("score", "text")]
    assert await reader.list_columns("Survey 2026", "Answers") == columns
    assert await read(reader, table("Survey 2026", "Answers")) == (
        columns,
        [("An", "7.5"), ("Binh", "8")],
    )
    assert await reason(read(reader, table("Survey 2026", "Gone"))) == (
        "source_not_found",
        google_drive.NO_TAB,
    )
    assert await reason(read(reader, table("Private", "Secret"))) == ("source_not_found", NO_FILE)
    assert drive.downloads == []
    assert {path.split("/")[3] for path, _ in drive.sheets.requests} == {SHEET_ID}
    assert oauth.refreshed == [REFRESH_TOKEN]


@pytest.mark.parametrize(
    ("answer", "expected"),
    [
        (401, "access_revoked"),
        (403, "permission_denied"),
        (404, "permission_denied"),
        (429, "rate_limited"),
        (500, "unreachable"),
        (httpx.ReadTimeout("slow"), "timeout"),
        (httpx.ConnectError("down"), "unreachable"),
        (httpx.Response(200, text="<html>not json</html>"), "unreachable"),
        (httpx.Response(200, json={"files": [{"id": "x"}]}), "unreachable"),
        # A redirect is not followed: the token goes to Google's address only.
        (httpx.Response(302, headers={"location": "https://evil.example/"}), "unreachable"),
    ],
)
async def test_what_drive_answers_decides_the_reason(executor, answer, expected) -> None:
    drive, oauth = folder(**{"f-csv": csv_file("scores.csv")})
    drive.fail_with = answer
    reader = connector(drive, oauth, executor)
    assert (await reason(reader.list_schemas()))[0] == expected
    assert (await reason(read(reader, table("scores.csv"))))[0] == expected
    if not isinstance(answer, httpx.Response) or answer.status_code != 200:
        assert (await reason(reader.test()))[0] == expected


async def test_a_search_drive_rejects_is_not_a_file_that_is_missing(executor) -> None:
    drive, oauth = folder(**{"f-csv": csv_file("scores.csv")})
    drive.fail_with = 400
    reader = connector(drive, oauth, executor)
    for reading in (
        reader.list_schemas(),
        reader.list_tables("scores.csv", search=None, limit=10),
        read(reader, table("scores.csv")),
    ):
        assert (await reason(reading))[0] == "unreachable"


@pytest.mark.parametrize(
    ("answer", "expected"),
    [
        (401, "access_revoked"),
        (403, "permission_denied"),
        (404, "permission_denied"),
        (429, "rate_limited"),
        (500, "unreachable"),
        (httpx.ReadTimeout("slow"), "timeout"),
        (httpx.Response(302, headers={"location": "https://evil.example/"}), "unreachable"),
    ],
)
async def test_a_download_that_fails_leaves_no_file(
    executor, temporary_files, answer, expected
) -> None:
    drive, oauth = folder(**{"f-csv": csv_file("scores.csv")})
    reader = connector(drive, oauth, executor)
    listing = drive._listing

    def fail_from_here(params):
        # The file is found, and Drive then refuses to hand it over.
        drive.fail_with = answer
        return listing(params)

    drive._listing = fail_from_here
    assert (await reason(read(reader, table("scores.csv"))))[0] == expected
    assert len(temporary_files) == 1 and temporary_files[0].closed


async def test_a_read_cancelled_while_the_file_arrives_leaves_no_file(
    executor, temporary_files, monkeypatch
) -> None:
    monkeypatch.setattr(google_drive, "ROWS_PER_BATCH", 1)
    drive, oauth = folder(**{"f-csv": csv_file("scores.csv")})
    drive.stall_downloads = asyncio.Event()
    reader = connector(drive, oauth, executor)

    reading = asyncio.create_task(read(reader, table("scores.csv")))
    await asyncio.wait_for(drive.stalled.wait(), timeout=5)
    (handle,) = temporary_files
    assert not handle.closed
    reading.cancel()
    with pytest.raises(asyncio.CancelledError):
        await reading
    assert handle.closed

    # Left part-way through the rows: the file goes all the same.
    drive.stall_downloads = None
    async with reader.open_rows(table("scores.csv"), max_rows=None) as stream:
        assert to_text((await anext(stream.rows))[0]) == "An"
        assert not temporary_files[-1].closed
    assert temporary_files[-1].closed


async def test_the_token_is_sent_to_google_only(executor) -> None:
    drive, oauth = folder(**{"f-csv": csv_file("scores.csv")})
    seen = []

    async def watching(request: httpx.Request) -> httpx.Response:
        seen.append((request.url.host, request.headers["authorization"]))
        return await drive._answer(request)

    reader = connector(drive, oauth, executor)
    reader._account._transport = httpx.MockTransport(watching)
    await reader.test()
    await read(reader, table("scores.csv"))
    assert set(seen) == {("www.googleapis.com", f"Bearer {access_token_for(REFRESH_TOKEN)}")}
